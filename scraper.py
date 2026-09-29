"""
Söker efter jobbinlägg om "statist" / "skådespelare" på LinkedIn, X/Twitter,
Instagram och Facebook, ENDAST från den senaste månaden, och exporterar
resultaten till CSV + Excel med samma kolumner som det gamla Google-arket:
Platform | Title | URL | Snippet | Date

Sökbackends som stöds (välj med SEARCH_BACKEND i .env-filen):
  - serper  -> https://serper.dev (Google-resultat, 2 500 gratis sökningar)
  - google  -> Google Custom Search JSON API (endast om du redan har en nyckel:
               stängd för nya kunder, slutar fungera 2027-01-01)
"""

import csv
import os
import re
import sys
import time
from datetime import datetime

import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font

load_dotenv()

# Windows-terminaler använder ibland cp1252: tvinga UTF-8 så att å ä ö och
# emojis inte kraschar utskriften (t.ex. vid omdirigering till en loggfil).
for stream in (sys.stdout, sys.stderr):
    stream.reconfigure(encoding="utf-8", errors="replace")

# ---------------------------------------------------------------- INSTÄLLNINGAR
BACKEND = os.getenv("SEARCH_BACKEND", "serper").strip().lower()
PAGES_PER_QUERY = int(os.getenv("PAGES_PER_QUERY", "3"))  # 10 resultat/sida

KEYWORDS = [
    "statist",
    "statister",
    "skådespelare",
    "casting",
    "statistjobb",
]

PLATFORMS = {
    "LinkedIn": "site:linkedin.com",
    "Twitter/X": "(site:twitter.com OR site:x.com)",
    "Instagram": "site:instagram.com",
    "Facebook": "site:facebook.com",
}

# Ett resultat behålls bara om titeln eller utdraget matchar minst ett av
# mönstren. Hela ord, så att "statist" INTE träffar "statistik"/"statistiker".
# Lämna listan tom för att inte filtrera alls.
MUST_MATCH_ANY = [
    r"\bstatist(er|erna|en|jobb\w*)?\b",
    r"\bskådespel\w*",
    r"\bcasting\w*",
    r"\bsöker\b",
    r"\bsökes\b",
    r"\bextras?\b",
]
# ------------------------------------------------------------------------------

RELEVANT_RE = [re.compile(p, re.IGNORECASE) for p in MUST_MATCH_ANY]


def search_serper(query: str, page: int) -> list[dict]:
    resp = requests.post(
        "https://google.serper.dev/search",
        headers={"X-API-KEY": os.environ["SERPER_API_KEY"],
                 "Content-Type": "application/json"},
        json={
            "q": query,
            "gl": "se",          # resultat från Sverige
            "hl": "sv",          # svenska
            "tbs": "qdr:m",      # <-- ENDAST SENASTE MÅNADEN
            "num": 10,
            "page": page,
        },
        timeout=30,
    )
    resp.raise_for_status()
    return [
        {"title": r.get("title", ""), "url": r.get("link", ""),
         "snippet": r.get("snippet", ""), "date": r.get("date", "")}
        for r in resp.json().get("organic", [])
    ]


def search_google_cse(query: str, page: int) -> list[dict]:
    resp = requests.get(
        "https://www.googleapis.com/customsearch/v1",
        params={
            "key": os.environ["GOOGLE_API_KEY"],
            "cx": os.environ["GOOGLE_CX"],
            "q": query,
            "gl": "se",
            "lr": "lang_sv",
            "dateRestrict": "m1",   # <-- ENDAST SENASTE MÅNADEN
            "num": 10,
            "start": (page - 1) * 10 + 1,
        },
        timeout=30,
    )
    resp.raise_for_status()
    return [
        {"title": r.get("title", ""), "url": r.get("link", ""),
         "snippet": r.get("snippet", "").replace("\n", " "), "date": ""}
        for r in resp.json().get("items", [])
    ]


BACKENDS = {
    "serper": (search_serper, ["SERPER_API_KEY"]),
    "google": (search_google_cse, ["GOOGLE_API_KEY", "GOOGLE_CX"]),
}


def check_config():
    """Stoppar direkt med ett begripligt meddelande om .env är fel ifylld."""
    if BACKEND not in BACKENDS:
        sys.exit(f"❌ Okänd SEARCH_BACKEND '{BACKEND}'. "
                 f"Tillåtna värden: {', '.join(BACKENDS)}.")
    search, required = BACKENDS[BACKEND]
    missing = [k for k in required if not os.getenv(k, "").strip()]
    if missing:
        sys.exit(f"❌ Saknar {', '.join(missing)} i .env-filen "
                 f"(kopiera env.example till .env och fyll i nyckeln).")
    return search


def is_relevant(item: dict) -> bool:
    if not RELEVANT_RE:
        return True
    text = f"{item['title']} {item['snippet']}"
    return any(p.search(text) for p in RELEVANT_RE)


def collect(search, rows: list[dict]) -> None:
    """Fyller på rows under körningen, så att inget går förlorat vid avbrott."""
    keyword_block = " OR ".join(f'"{k}"' for k in KEYWORDS)
    seen_urls: set[str] = set()

    for platform, site_filter in PLATFORMS.items():
        query = f"{site_filter} ({keyword_block})"
        print(f"\n🔎 {platform}: {query}")
        for page in range(1, PAGES_PER_QUERY + 1):
            try:
                results = search(query, page)
            except requests.RequestException as e:
                # Nätverksfel eller API-fel: hoppa till nästa plattform,
                # men behåll allt som redan har samlats in.
                print(f"   ⚠️  fel på sida {page}: {e}")
                break
            if not results:
                break
            new = 0
            for r in results:
                if r["url"] in seen_urls or not is_relevant(r):
                    continue
                seen_urls.add(r["url"])
                rows.append({"Platform": platform, "Title": r["title"],
                             "URL": r["url"], "Snippet": r["snippet"],
                             "Date": r["date"]})
                new += 1
            print(f"   sida {page}: {len(results)} resultat, {new} nya")
            time.sleep(1)  # var snäll mot API:et


def export(rows: list[dict]) -> None:
    stamp = datetime.now().strftime("%Y-%m-%d")
    os.makedirs("output", exist_ok=True)
    cols = ["Platform", "Title", "URL", "Snippet", "Date"]

    # CSV (utf-8-sig så att Excel visar å ä ö korrekt)
    csv_path = f"output/social_media_inlagg_{stamp}.csv"
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)

    # Excel (kan importeras direkt i Google Kalkylark)
    xlsx_path = f"output/social_media_inlagg_{stamp}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Resultat"
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in rows:
        ws.append([row[c] for c in cols])
        link_cell = ws.cell(row=ws.max_row, column=3)
        link_cell.hyperlink = row["URL"]
        link_cell.style = "Hyperlink"
    for col, width in zip("ABCDE", (12, 60, 60, 100, 14)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    wb.save(xlsx_path)

    print(f"\n✅ {len(rows)} resultat sparade i:\n   {csv_path}\n   {xlsx_path}")


def main() -> None:
    search = check_config()
    rows: list[dict] = []
    try:
        collect(search, rows)
    except KeyboardInterrupt:
        print("\n⏹  Avbruten – sparar det som hunnit samlas in.")
    if not rows:
        print("\nInga resultat hittades.")
        return
    export(rows)


if __name__ == "__main__":
    main()
